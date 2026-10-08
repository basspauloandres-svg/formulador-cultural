from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any
import math
import re
import textwrap
from xml.sax.saxutils import escape as xml_escape

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    KeepTogether,
)

DOC_FOOTER = "Formulador Cultural · Desarrollo por Paulo Olarte"
POR_VERIFICAR = "[POR VERIFICAR]"


@dataclass(frozen=True)
class DocumentTheme:
    font: str = "Aptos"
    body_size: int = 10
    title_size: int = 22
    heading1_size: int = 15
    heading2_size: int = 12
    line_spacing: float = 1.28
    margin_top_cm: float = 2.3
    margin_bottom_cm: float = 2.0
    margin_left_cm: float = 2.5
    margin_right_cm: float = 2.5
    accent_hex: str = "176B49"
    muted_hex: str = "66717C"
    line_hex: str = "DCE7E1"


THEME = DocumentTheme()


def clean(value: Any) -> str:
    text = " ".join(str(value or "").split())
    return text or POR_VERIFICAR


def rows(value: Any) -> list[dict[str, Any]]:
    return value if isinstance(value, list) else []


def _sentence_chunks(value: Any, max_sentences: int = 4) -> list[str]:
    """Split long prose into readable paragraphs without altering wording."""
    text = clean(value)
    if text == POR_VERIFICAR:
        return [text]
    parts = [x.strip() for x in re.split(r"(?<=[.!?])\s+(?=[A-ZÁÉÍÓÚÜÑ0-9¿¡\[])", text) if x.strip()]
    if len(parts) <= max_sentences:
        return [text]
    return [" ".join(parts[i:i + max_sentences]) for i in range(0, len(parts), max_sentences)]


def _compact_text(value: Any, max_chars: int = 120) -> str:
    text = clean(value)
    return text if len(text) <= max_chars else text[:max_chars - 1].rstrip(" ,.;:") + "…"


def normalized_payload(payload: dict[str, Any]) -> dict[str, Any]:
    """Accepts frontend-export-like structured data and preserves unknown values."""
    return {
        "title": clean(payload.get("title") or payload.get("project_name")),
        "entity": clean(payload.get("entity")),
        "territory": clean(payload.get("territory")),
        "responsible": clean(payload.get("responsible")),
        "summary": clean(payload.get("summary")),
        "context": clean(payload.get("context")),
        "population": clean(payload.get("population")),
        "evidence": clean(payload.get("evidence")),
        "sources": clean(payload.get("sources")),
        "problem": clean(payload.get("problem")),
        "strategy": clean(payload.get("strategy")),
        "objective_general": clean(payload.get("objective_general")),
        "objectives": rows(payload.get("objectives")),
        "results": rows(payload.get("results")),
        "activities": rows(payload.get("activities")),
        "indicators": rows(payload.get("indicators")),
        "schedule": rows(payload.get("schedule")),
        "budget": rows(payload.get("budget")),
        "risks": rows(payload.get("risks")),
        "problem_tree": rows(payload.get("problem_tree")),
        "objective_tree": rows(payload.get("objective_tree")),
        "vester": payload.get("vester") if isinstance(payload.get("vester"), dict) else {},
        "coherence": payload.get("coherence") if isinstance(payload.get("coherence"), dict) else {},
    }


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def _configure_docx_styles(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Cm(THEME.margin_top_cm)
    sec.bottom_margin = Cm(THEME.margin_bottom_cm)
    sec.left_margin = Cm(THEME.margin_left_cm)
    sec.right_margin = Cm(THEME.margin_right_cm)

    normal = doc.styles["Normal"]
    normal.font.name = THEME.font
    normal.font.size = Pt(THEME.body_size)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = THEME.line_spacing

    h1 = doc.styles["Heading 1"]
    h1.font.name = THEME.font
    h1.font.size = Pt(THEME.heading1_size)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(THEME.accent_hex)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = THEME.font
    h2.font.size = Pt(THEME.heading2_size)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string("244A37")
    h2.paragraph_format.keep_with_next = True


def _add_docx_footer(doc: Document) -> None:
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(DOC_FOOTER)
        r.font.name = THEME.font
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(THEME.muted_hex)


def _add_docx_page_number(section) -> None:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(" · Página ")
    run = footer.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def _add_docx_toc(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "Actualiza la tabla de contenido al abrir el documento."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r._r.extend([fld_char, instr, fld_sep, fallback, fld_end])


def _docx_cover(doc: Document, p: dict[str, Any]) -> None:
    for _ in range(4):
        doc.add_paragraph()
    eyebrow = doc.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = eyebrow.add_run("FORMATO ESTÁNDAR DE PROYECTO CULTURAL")
    rr.bold = True
    rr.font.size = Pt(10)
    rr.font.color.rgb = RGBColor.from_string(THEME.accent_hex)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(p["title"])
    r.bold = True
    r.font.name = THEME.font
    r.font.size = Pt(THEME.title_size)
    r.font.color.rgb = RGBColor.from_string("173D2C")

    for label, value in (
        ("Entidad / organización", p["entity"]),
        ("Territorio", p["territory"]),
        ("Responsable", p["responsible"]),
    ):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(f"{label}: {value}")
        run.font.size = Pt(10)

    doc.add_paragraph()
    stamp = doc.add_paragraph()
    stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = stamp.add_run(DOC_FOOTER)
    rs.italic = True
    rs.font.size = Pt(8)
    rs.font.color.rgb = RGBColor.from_string(THEME.muted_hex)
    doc.add_page_break()


def _docx_body_paragraph(doc: Document, text: str, bold_label: str | None = None) -> None:
    chunks = _sentence_chunks(text, 4)
    for index, chunk in enumerate(chunks):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.28
        p.paragraph_format.space_after = Pt(9)
        if bold_label and index == 0:
            rb = p.add_run(bold_label)
            rb.bold = True
        p.add_run(chunk)
        if (index + 1) % 4 == 0 and index < len(chunks) - 1:
            doc.add_page_break()


def _docx_table(doc: Document, headers: list[str], body: list[list[Any]], widths_cm: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    _set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        _shade_cell(cell, "EEF5F1")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = THEME.font
            run.font.size = Pt(8)
        _set_cell_margins(cell)
    for row in body:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = clean(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for par in cells[i].paragraphs:
                for run in par.runs:
                    run.font.name = THEME.font
                    run.font.size = Pt(8)
            _set_cell_margins(cells[i])
    if widths_cm:
        for row in table.rows:
            for i, w in enumerate(widths_cm):
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)
    doc.add_paragraph()


def _docx_indicator_cards(doc: Document, indicators: list[dict[str, Any]]) -> None:
    fields = [
        ("Indicador", "indicator"),
        ("Elemento asociado", "linkedText"),
        ("Fórmula / criterio", "formula"),
        ("Unidad", "unidad"),
        ("Línea base", "lineaBase"),
        ("Meta", "meta"),
        ("Periodicidad", "periodicidad"),
        ("Medio de verificación", "medioVerificacion"),
        ("Responsable", "responsable"),
        ("Plazo", "plazo"),
    ]
    for n, item in enumerate(indicators, 1):
        level = clean(item.get("linkedType") or "Indicador")
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(4)
        run = h.add_run(f"Indicador {n} · {level}")
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("244A37")
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for label, key in fields:
            row = table.add_row().cells
            row[0].text = label
            value = item.get(key)
            if key == "linkedText" and not value:
                value = item.get("activityText")
            row[1].text = clean(value)
            _shade_cell(row[0], "EEF5F1")
            for cell in row:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                _set_cell_margins(cell, top=80, start=90, bottom=80, end=90)
                for par in cell.paragraphs:
                    par.paragraph_format.line_spacing = 1.12
                    for rr in par.runs:
                        rr.font.name = THEME.font
                        rr.font.size = Pt(8.5)
            for rr in row[0].paragraphs[0].runs:
                rr.bold = True
        doc.add_paragraph()


def _docx_resource_entries(doc: Document, budget: list[dict[str, Any]]) -> None:
    if not budget:
        _docx_body_paragraph(doc, POR_VERIFICAR)
        return
    for n, item in enumerate(budget, 1):
        activity = clean(item.get("activityId") or item.get("activityText") or item.get("activity"))
        description = clean(item.get("description"))
        source = clean(item.get("fundingSource"))
        line = (
            f"{activity} — Recurso: {description}; Unidad: {clean(item.get('unit'))}; "
            f"Cantidad: {clean(item.get('quantity'))}; Veces: {clean(item.get('frequency'))}; "
            f"Costo unitario: {clean(item.get('unitCost'))}; Total: {clean(item.get('totalCost'))}; "
            f"Tipo: {clean(item.get('costType'))}; Fuente: {source}."
        )
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        p.paragraph_format.line_spacing = 1.22
        p.paragraph_format.space_after = Pt(7)
        r = p.add_run(f"{n}. ")
        r.bold = True
        p.add_run(line)


def _save_figure(fig, target: Path) -> Path:
    fig.savefig(target, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return target


def _wrap_node_text(value: Any, width: int = 38, max_lines: int = 5) -> str:
    text = clean(value)
    lines = textwrap.wrap(text, width=max(18, width), break_long_words=False, break_on_hyphens=False)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        if lines:
            lines[-1] = lines[-1].rstrip(" .") + "…"
    return "\n".join(lines) or POR_VERIFICAR


def _tree_figure(nodes: list[dict[str, Any]], title: str, target: Path) -> Path | None:
    """Render a legible tree with adaptive boxes and wrapped text.

    The document uses this raster figure in both DOCX and PDF, so the layout must
    remain readable at A4 width and on mobile document viewers.
    """
    if not nodes:
        return None
    order = ["indirect_effect", "direct_effect", "central", "direct_cause", "indirect_cause"]
    labels = {
        "indirect_effect": "Efectos / fines indirectos",
        "direct_effect": "Efectos / fines directos",
        "central": "Problema / objetivo central",
        "direct_cause": "Causas / medios directos",
        "indirect_cause": "Causas / medios indirectos",
    }
    active_layers = [(zone, [n for n in nodes if n.get("zone") == zone]) for zone in order]
    active_layers = [(zone, layer) for zone, layer in active_layers if layer]
    max_in_layer = max((len(layer) for _, layer in active_layers), default=1)
    fig_h = max(5.4, 1.55 * len(active_layers) + 1.0)
    fig_w = max(11.5, 3.8 * min(max_in_layer, 3))
    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    ax.set_axis_off()
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)

    layout: dict[str, dict[str, Any]] = {}
    top, bottom = .88, .10
    layer_gap = (top - bottom) / max(1, len(active_layers) - 1) if len(active_layers) > 1 else 0

    for li, (zone, layer) in enumerate(active_layers):
        y = top - li * layer_gap if len(active_layers) > 1 else .5
        count = len(layer)
        horizontal_margin = .055
        gap = .035 if count <= 3 else .02
        available = 1 - 2 * horizontal_margin - gap * max(0, count - 1)
        box_w = min(.46 if zone == "central" else .36, available / max(1, count))
        total = box_w * count + gap * max(0, count - 1)
        start = .5 - total / 2
        wrap_width = max(24, int(52 * box_w / .36))
        for i, n in enumerate(layer):
            x = start + box_w / 2 + i * (box_w + gap)
            wrapped = _wrap_node_text(n.get("text"), width=wrap_width, max_lines=5)
            line_count = max(1, wrapped.count("\n") + 1)
            box_h = min(.16, .055 + .020 * line_count)
            nid = str(n.get("id", f"{zone}-{i}"))
            layout[nid] = {
                "x": x, "y": y, "w": box_w, "h": box_h, "zone": zone,
                "text": wrapped, "label": labels[zone],
            }

    central = next((n for n in nodes if n.get("zone") == "central"), None)
    central_id = str(central.get("id")) if central else None

    # Connectors first so lines never cross over node text.
    for n in nodes:
        nid = str(n.get("id"))
        src = layout.get(nid)
        if not src or n.get("zone") == "central":
            continue
        parent_id = str(n.get("parentId")) if n.get("parentId") else central_id
        dst = layout.get(parent_id) if parent_id else None
        if not dst:
            continue
        src_y = src["y"] + (src["h"] / 2 if dst["y"] > src["y"] else -src["h"] / 2)
        dst_y = dst["y"] - (dst["h"] / 2 if dst["y"] > src["y"] else -dst["h"] / 2)
        ax.plot([src["x"], src["x"], dst["x"]], [src_y, (src_y + dst_y) / 2, dst_y],
                color="#93A19A", linewidth=1.05, zorder=1)

    for item in layout.values():
        central_zone = item["zone"] == "central"
        x, y, w, h = item["x"], item["y"], item["w"], item["h"]
        box = FancyBboxPatch(
            (x - w / 2, y - h / 2), w, h,
            boxstyle="round,pad=0.008,rounding_size=0.012",
            linewidth=1.5,
            edgecolor="#B27A00" if central_zone else "#2D8A61",
            facecolor="#FFF7D6" if central_zone else "#F5FBF7",
            zorder=2,
        )
        ax.add_patch(box)
        ax.text(x, y, item["text"], ha="center", va="center", fontsize=8.2,
                color="#23312B", linespacing=1.25, zorder=3)
        ax.text(x, y + h / 2 + .018, item["label"], ha="center", va="bottom",
                fontsize=7.2, color="#66717C", zorder=3)

    ax.text(.5, .97, title, ha="center", va="top", fontsize=13, fontweight="bold", color="#173D2C")
    fig.subplots_adjust(left=.02, right=.98, top=.98, bottom=.02)
    return _save_figure(fig, target)


def _vester_figure(vester: dict[str, Any], target: Path) -> Path | None:
    rows_ = rows(vester.get("rows"))
    if not rows_:
        return None
    xs = [float(r.get("influence", 0)) for r in rows_]
    ys = [float(r.get("dependence", 0)) for r in rows_]
    mx = float(vester.get("meanInfluence", sum(xs) / max(1, len(xs))))
    my = float(vester.get("meanDependence", sum(ys) / max(1, len(ys))))
    fig, ax = plt.subplots(figsize=(8.5, 5.6))
    ax.axvline(mx, color="#9aa5ae", ls="--", lw=1)
    ax.axhline(my, color="#9aa5ae", ls="--", lw=1)
    ax.scatter(xs, ys, s=55)
    for i, (x, y) in enumerate(zip(xs, ys), 1):
        ax.annotate(f"P{i}", (x, y), xytext=(5, 5), textcoords="offset points", fontsize=8)
    ax.set_xlabel("Influencia")
    ax.set_ylabel("Dependencia")
    ax.set_title("Matriz Vester · plano de influencia y dependencia")
    ax.grid(alpha=.15)
    return _save_figure(fig, target)


def _gantt_figure(schedule: list[dict[str, Any]], target: Path) -> Path | None:
    import datetime as dt
    valid = []
    for row in schedule:
        try:
            start = dt.datetime.fromisoformat(str(row.get("startDate") or row.get("start")))
            end = dt.datetime.fromisoformat(str(row.get("endDate") or row.get("end")))
            valid.append((clean(row.get("activityText") or row.get("activity")), start, end))
        except Exception:
            continue
    if not valid:
        return None
    fig_h = max(3.8, 0.48 * len(valid) + 1.6)
    fig, ax = plt.subplots(figsize=(10.5, fig_h))
    for i, (name, start, end) in enumerate(valid):
        width = max((end - start).days, 1)
        ax.barh(i, width, left=start.toordinal(), height=.55)
    ax.set_yticks(range(len(valid)))
    ax.set_yticklabels([_compact_text(v[0], 46) for v in valid], fontsize=7.5)
    ticks = ax.get_xticks()
    ax.set_xticklabels([dt.date.fromordinal(int(t)).isoformat() if t > 1 else "" for t in ticks], rotation=25, ha="right", fontsize=7)
    ax.set_title("Cronograma gráfico")
    ax.grid(axis="x", alpha=.15)
    fig.tight_layout()
    return _save_figure(fig, target)


def _budget_figure(budget: list[dict[str, Any]], target: Path) -> Path | None:
    sums: dict[str, float] = {}
    for row in budget:
        value = row.get("totalCost", row.get("total", 0))
        try:
            amount = float(value)
        except Exception:
            continue
        label = clean(row.get("activityText") or row.get("activity") or row.get("description"))
        sums[label] = sums.get(label, 0.0) + amount
    if not sums:
        return None
    items = sorted(sums.items(), key=lambda x: x[1], reverse=True)[:10]
    fig_h = max(3.5, .42 * len(items) + 1.4)
    fig, ax = plt.subplots(figsize=(9.5, fig_h))
    ax.barh(range(len(items)), [v for _, v in items])
    ax.set_yticks(range(len(items)))
    ax.set_yticklabels([_compact_text(k, 44) for k, _ in items], fontsize=7.5)
    ax.invert_yaxis()
    ax.set_title("Distribución del presupuesto por actividad")
    ax.set_xlabel("Valor")
    ax.grid(axis="x", alpha=.15)
    fig.tight_layout()
    return _save_figure(fig, target)


def _docx_add_figure(
    doc: Document,
    path: Path | None,
    caption: str,
    width_inches: float = 6.35,
    standalone: bool = False,
) -> None:
    if not path or not path.exists():
        return
    if standalone:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width_inches))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(8)
    cap.runs[0].font.color.rgb = RGBColor.from_string(THEME.muted_hex)
    if standalone:
        doc.add_page_break()


def _indicator_summary_rows(indicators: list[dict[str, Any]]) -> list[list[Any]]:
    return [[
        x.get("linkedType", "Actividad"),
        x.get("linkedText") or x.get("activityText"),
        x.get("indicator"),
        x.get("lineaBase"),
        x.get("meta"),
        x.get("medioVerificacion"),
    ] for x in indicators]


def _indicator_technical_rows(indicators: list[dict[str, Any]]) -> list[list[Any]]:
    return [[
        x.get("linkedType", "Actividad"),
        x.get("indicator"),
        x.get("formula"),
        x.get("unidad"),
        x.get("lineaBase"),
        x.get("meta"),
        x.get("periodicidad"),
        x.get("medioVerificacion"),
        x.get("responsable"),
        x.get("plazo"),
    ] for x in indicators]


def generate_docx(payload: dict[str, Any]) -> bytes:
    p = normalized_payload(payload)
    doc = Document()
    _configure_docx_styles(doc)
    _add_docx_footer(doc)
    _add_docx_page_number(doc.sections[0])
    _docx_cover(doc, p)

    doc.add_heading("Tabla de contenido", level=1)
    _add_docx_toc(doc)
    doc.add_page_break()

    with TemporaryDirectory() as td:
        td = Path(td)
        figures = {
            "vester": _vester_figure(p["vester"], td / "vester.png"),
            "problem": _tree_figure(p["problem_tree"], "Árbol de problemas", td / "problem_tree.png"),
            "objective": _tree_figure(p["objective_tree"], "Árbol de objetivos", td / "objective_tree.png"),
            "gantt": _gantt_figure(p["schedule"], td / "gantt.png"),
            "budget": _budget_figure(p["budget"], td / "budget.png"),
        }

        sections: list[tuple[str, callable]] = [
            ("1. Resumen ejecutivo", lambda: _docx_body_paragraph(doc, p["summary"])),
            ("2. Contexto territorial, cultural y social", lambda: _docx_body_paragraph(doc, p["context"])),
            ("3. Población", lambda: _docx_body_paragraph(doc, p["population"])),
            ("4. Evidencia y antecedentes", lambda: (_docx_body_paragraph(doc, p["evidence"]), _docx_body_paragraph(doc, p["sources"], "Fuentes: "))),
            ("5. Priorización de situaciones · Matriz Vester", lambda: _docx_add_figure(doc, figures["vester"], "Figura 1. Plano de influencia y dependencia de la matriz Vester.")),
            ("6. Planteamiento del problema", lambda: (_docx_body_paragraph(doc, p["problem"]), _docx_add_figure(doc, figures["problem"], "Figura 2. Árbol de problemas.", standalone=True))),
            ("7. Objetivos", lambda: (
                _docx_body_paragraph(doc, p["objective_general"], "Objetivo general: "),
                _docx_table(doc, ["Objetivos específicos"], [[o.get("text", o)] for o in p["objectives"]]) if p["objectives"] else None,
                _docx_add_figure(doc, figures["objective"], "Figura 3. Árbol de objetivos.", standalone=True)
            )),
            ("8. Estrategia de intervención", lambda: _docx_body_paragraph(doc, p["strategy"])),
            ("9. Resultados esperados", lambda: _docx_table(doc, ["Resultados"], [[r.get("text", r)] for r in p["results"]]) if p["results"] else _docx_body_paragraph(doc, POR_VERIFICAR)),
            ("10. Plan de actividades", lambda: _docx_table(doc, ["ID", "Resultado / objetivo", "Actividad"], [[a.get("id"), a.get("resultText") or a.get("objectiveText"), a.get("text")] for a in p["activities"]], [2.2, 6.2, 9.0])),
            ("11. Indicadores y metas", lambda: (
                doc.add_heading("11.1 Tabla resumida", level=2),
                _docx_table(doc, ["Nivel", "Elemento", "Indicador", "Línea base", "Meta", "Medio de verificación"], _indicator_summary_rows(p["indicators"]), [1.7, 4.0, 4.5, 2.0, 2.0, 3.2]),
                doc.add_heading("11.2 Fichas técnicas de indicadores", level=2),
                _docx_indicator_cards(doc, p["indicators"])
            )),
            ("12. Cronograma", lambda: (
                _docx_add_figure(doc, figures["gantt"], "Figura 4. Cronograma gráfico tipo Gantt."),
                _docx_table(doc, ["Actividad", "Inicio", "Fin", "Responsable", "Frecuencia"], [[s.get("activityText"), s.get("startDate"), s.get("endDate"), s.get("responsible"), s.get("frequency")] for s in p["schedule"]], [8.2, 2.5, 2.5, 3.4, 2.4])
            )),
            ("13. Recursos y presupuesto", lambda: (
                _docx_add_figure(doc, figures["budget"], "Figura 5. Distribución del presupuesto por actividad."),
                doc.add_heading("13.1 Soportes de recursos", level=2),
                _docx_resource_entries(doc, p["budget"])
            )),
            ("14. Riesgos y respuestas", lambda: _docx_table(doc, ["Origen", "Riesgo", "Probabilidad", "Impacto", "Nivel", "Prevención", "Contingencia", "Responsable"], [[r.get("linkedObjectType"), r.get("event"), r.get("probability"), r.get("impact"), r.get("riskLevel"), r.get("preventiveResponse"), r.get("contingencyResponse"), r.get("owner")] for r in p["risks"]])),
            ("15. Coherencia y trazabilidad", lambda: _docx_body_paragraph(doc, clean(p["coherence"].get("summary") or f"Índice orientativo: {p['coherence'].get('score', POR_VERIFICAR)}"))),
            ("16. Fuentes y anexos", lambda: _docx_body_paragraph(doc, p["sources"])),
        ]

        for title, builder in sections:
            doc.add_heading(title, level=1)
            builder()

        out = BytesIO()
        doc.save(out)
        return out.getvalue()


class _NumberedCanvasMixin:
    pass


def _pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="FC_Title", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=22, leading=26, alignment=TA_CENTER, textColor=colors.HexColor("#173D2C"),
        spaceAfter=12
    ))
    styles.add(ParagraphStyle(
        name="FC_H1", parent=styles["Heading1"], fontName="Helvetica-Bold",
        fontSize=15, leading=18, textColor=colors.HexColor("#176B49"),
        spaceBefore=10, spaceAfter=7
    ))
    styles.add(ParagraphStyle(
        name="FC_H2", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=11, leading=14, textColor=colors.HexColor("#244A37"),
        spaceBefore=8, spaceAfter=5
    ))
    styles.add(ParagraphStyle(
        name="FC_Body", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=9.5, leading=15.0, alignment=TA_JUSTIFY, spaceAfter=8
    ))
    styles.add(ParagraphStyle(
        name="FC_Table", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=7.2, leading=9.2, alignment=TA_LEFT, spaceAfter=0
    ))
    styles.add(ParagraphStyle(
        name="FC_TableHead", parent=styles["BodyText"], fontName="Helvetica-Bold",
        fontSize=7.2, leading=9.2, alignment=TA_LEFT, textColor=colors.HexColor("#244A37"), spaceAfter=0
    ))
    styles.add(ParagraphStyle(
        name="FC_Record", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=8.8, leading=12.0, leftIndent=12, firstLineIndent=-12, spaceAfter=7
    ))
    styles.add(ParagraphStyle(
        name="FC_Small", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=7.8, leading=10, textColor=colors.HexColor("#66717C")
    ))
    return styles


def _pdf_table(headers: list[str], body: list[list[Any]], widths=None):
    styles = _pdf_styles()
    head_style = styles["FC_TableHead"]
    cell_style = styles["FC_Table"]
    data = [[Paragraph(xml_escape(clean(h)), head_style) for h in headers]]
    for row in body:
        data.append([Paragraph(xml_escape(clean(v)), cell_style) for v in row])
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="CENTER", splitByRow=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF5F1")),
        ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#DCE7E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def _pdf_indicator_cards(story: list[Any], indicators: list[dict[str, Any]], styles) -> None:
    fields = [
        ("Indicador", "indicator"),
        ("Elemento asociado", "linkedText"),
        ("Fórmula / criterio", "formula"),
        ("Unidad", "unidad"),
        ("Línea base", "lineaBase"),
        ("Meta", "meta"),
        ("Periodicidad", "periodicidad"),
        ("Medio de verificación", "medioVerificacion"),
        ("Responsable", "responsable"),
        ("Plazo", "plazo"),
    ]
    for n, item in enumerate(indicators, 1):
        level = clean(item.get("linkedType") or "Indicador")
        story.append(Paragraph(f"<b>Indicador {n} · {xml_escape(level)}</b>", styles["FC_H2"]))
        body = []
        for label, key in fields:
            value = item.get(key)
            if key == "linkedText" and not value:
                value = item.get("activityText")
            body.append([label, value])
        story.append(_pdf_table(["Campo", "Información"], body, [40 * mm, 125 * mm]))
        story.append(Spacer(1, 4 * mm))


def _pdf_resource_entries(story: list[Any], budget: list[dict[str, Any]], styles) -> None:
    if not budget:
        story.append(Paragraph(POR_VERIFICAR, styles["FC_Body"]))
        return
    for n, item in enumerate(budget, 1):
        activity = clean(item.get("activityId") or item.get("activityText") or item.get("activity"))
        line = (
            f"<b>{n}. {xml_escape(activity)}</b> — "
            f"Recurso: {xml_escape(clean(item.get('description')))}; "
            f"Unidad: {xml_escape(clean(item.get('unit')))}; "
            f"Cantidad: {xml_escape(clean(item.get('quantity')))}; "
            f"Veces: {xml_escape(clean(item.get('frequency')))}; "
            f"Costo unitario: {xml_escape(clean(item.get('unitCost')))}; "
            f"Total: {xml_escape(clean(item.get('totalCost')))}; "
            f"Tipo: {xml_escape(clean(item.get('costType')))}; "
            f"Fuente: {xml_escape(clean(item.get('fundingSource')))}."
        )
        story.append(Paragraph(line, styles["FC_Record"]))


def generate_pdf(payload: dict[str, Any]) -> bytes:
    p = normalized_payload(payload)
    styles = _pdf_styles()
    out = BytesIO()

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#66717C"))
        footer = f"{DOC_FOOTER} · Página {doc.page}"
        canvas.drawCentredString(A4[0] / 2, 10 * mm, footer)
        canvas.restoreState()

    doc = BaseDocTemplate(
        out, pagesize=A4,
        rightMargin=16 * mm, leftMargin=16 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
        title=p["title"], author="Paulo Olarte"
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")
    doc.addPageTemplates([PageTemplate(id="main", frames=frame, onPage=on_page)])
    story = []

    with TemporaryDirectory() as td:
        td = Path(td)
        figs = {
            "vester": _vester_figure(p["vester"], td / "vester.png"),
            "problem": _tree_figure(p["problem_tree"], "Árbol de problemas", td / "problem_tree.png"),
            "objective": _tree_figure(p["objective_tree"], "Árbol de objetivos", td / "objective_tree.png"),
            "gantt": _gantt_figure(p["schedule"], td / "gantt.png"),
            "budget": _budget_figure(p["budget"], td / "budget.png"),
        }

        story += [Spacer(1, 42 * mm), Paragraph("FORMATO ESTÁNDAR DE PROYECTO CULTURAL", styles["FC_Small"]),
                  Paragraph(p["title"], styles["FC_Title"]),
                  Spacer(1, 6 * mm),
                  Paragraph(f"<b>Entidad / organización:</b> {p['entity']}", styles["FC_Body"]),
                  Paragraph(f"<b>Territorio:</b> {p['territory']}", styles["FC_Body"]),
                  Paragraph(f"<b>Responsable:</b> {p['responsible']}", styles["FC_Body"]),
                  Spacer(1, 20 * mm),
                  Paragraph(DOC_FOOTER, styles["FC_Small"]), PageBreak()]

        def add_h1(text): story.append(Paragraph(text, styles["FC_H1"]))
        def add_body(text):
            chunks = _sentence_chunks(text, 4)
            for index, chunk in enumerate(chunks):
                story.append(Paragraph(xml_escape(clean(chunk)), styles["FC_Body"]))
                if (index + 1) % 4 == 0 and index < len(chunks) - 1:
                    story.append(PageBreak())
        def add_img(path: Path | None, width_mm=160):
            if path and path.exists():
                img = Image(str(path), width=width_mm * mm)
                ratio = img.imageHeight / max(1, img.imageWidth)
                img.drawHeight = width_mm * mm * ratio
                story.append(img)
                story.append(Spacer(1, 3 * mm))

        add_h1("1. Resumen ejecutivo"); add_body(p["summary"])
        add_h1("2. Contexto territorial, cultural y social"); add_body(p["context"])
        add_h1("3. Población"); add_body(p["population"])
        add_h1("4. Evidencia y antecedentes"); add_body(p["evidence"]); add_body("Fuentes: " + p["sources"])
        story.append(PageBreak())
        add_h1("5. Priorización de situaciones · Matriz Vester"); add_img(figs["vester"])
        if rows(p["vester"].get("rows")):
            story.append(_pdf_table(["Situación", "Influencia", "Dependencia", "Clasificación"],
                                    [[r.get("text"), r.get("influence"), r.get("dependence"), r.get("quadrant")] for r in p["vester"]["rows"]],
                                    [85*mm, 25*mm, 25*mm, 28*mm]))
        story.append(PageBreak())
        add_h1("6. Planteamiento del problema"); add_body(p["problem"]); add_img(figs["problem"])
        story.append(PageBreak())
        add_h1("7. Objetivos"); add_body("Objetivo general: " + p["objective_general"])
        if p["objectives"]:
            story.append(_pdf_table(["Objetivos específicos"], [[o.get("text", o)] for o in p["objectives"]], [165*mm]))
        add_img(figs["objective"])
        add_h1("8. Estrategia de intervención"); add_body(p["strategy"])
        add_h1("9. Resultados esperados")
        if p["results"]:
            story.append(_pdf_table(["Resultados"], [[r.get("text", r)] for r in p["results"]], [165*mm]))
        add_h1("10. Plan de actividades")
        story.append(_pdf_table(["ID", "Resultado / objetivo", "Actividad"], [[a.get("id"), a.get("resultText") or a.get("objectiveText"), a.get("text")] for a in p["activities"]], [22*mm, 58*mm, 85*mm]))
        story.append(PageBreak())
        add_h1("11. Indicadores y metas")
        story.append(Paragraph("11.1 Tabla resumida", styles["FC_H2"]))
        story.append(_pdf_table(["Nivel","Elemento","Indicador","Línea base","Meta","Medio de verificación"], _indicator_summary_rows(p["indicators"]), [18*mm,35*mm,45*mm,18*mm,18*mm,31*mm]))
        story.append(Spacer(1, 4*mm))
        story.append(Paragraph("11.2 Fichas técnicas de indicadores", styles["FC_H2"]))
        _pdf_indicator_cards(story, p["indicators"], styles)
        story.append(PageBreak())
        add_h1("12. Cronograma"); add_img(figs["gantt"])
        story.append(_pdf_table(["Actividad","Inicio","Fin","Responsable","Frecuencia"], [[s.get("activityText"),s.get("startDate"),s.get("endDate"),s.get("responsible"),s.get("frequency")] for s in p["schedule"]], [75*mm,22*mm,22*mm,27*mm,20*mm]))
        story.append(PageBreak())
        add_h1("13. Recursos y presupuesto"); add_img(figs["budget"])
        story.append(Paragraph("13.1 Soportes de recursos", styles["FC_H2"]))
        _pdf_resource_entries(story, p["budget"], styles)
        add_h1("14. Riesgos y respuestas")
        story.append(_pdf_table(["Origen","Riesgo","Prob.","Impacto","Nivel","Prevención","Contingencia","Responsable"], [[r.get("linkedObjectType"),r.get("event"),r.get("probability"),r.get("impact"),r.get("riskLevel"),r.get("preventiveResponse"),r.get("contingencyResponse"),r.get("owner")] for r in p["risks"]], [18*mm,34*mm,14*mm,14*mm,14*mm,30*mm,30*mm,20*mm]))
        add_h1("15. Coherencia y trazabilidad"); add_body(clean(p["coherence"].get("summary") or f"Índice orientativo: {p['coherence'].get('score', POR_VERIFICAR)}"))
        add_h1("16. Fuentes y anexos"); add_body(p["sources"])

        doc.build(story)
    return out.getvalue()


def write_outputs(payload: dict[str, Any], output_dir: str | Path, stem: str = "proyecto-cultural") -> tuple[Path, Path]:
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    docx_path = out / f"{stem}.docx"
    pdf_path = out / f"{stem}.pdf"
    docx_path.write_bytes(generate_docx(payload))
    pdf_path.write_bytes(generate_pdf(payload))
    return docx_path, pdf_path
